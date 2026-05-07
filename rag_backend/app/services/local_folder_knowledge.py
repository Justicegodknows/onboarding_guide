import hashlib
import logging
from pathlib import Path
from typing import Any, Dict, List, Tuple

from app.core.config import settings

logger = logging.getLogger(__name__)

# Resolved once at import time: rag_backend/ → docs/EUZ Project/
_BACKEND_ROOT = Path(__file__).resolve().parents[2]

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".xls", ".txt"}


def _chunk_text(text: str, chunk_size: int, overlap: int = 100) -> List[str]:
    chunks = []
    start = 0
    while start < len(text):
        piece = text[start : start + chunk_size]
        if piece.strip():
            chunks.append(piece)
        start += chunk_size - overlap
    return chunks


def _extract_pdf(path: Path) -> str:
    import fitz  # pymupdf
    text = ""
    with fitz.open(str(path)) as doc:
        for page in doc:
            text += page.get_text()
    return text


def _extract_docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _extract_xlsx(path: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    lines = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            row_text = " | ".join(str(c) for c in row if c is not None and str(c).strip())
            if row_text.strip():
                lines.append(row_text)
    wb.close()
    return "\n".join(lines)


def _extract_xls(path: Path) -> str:
    import xlrd
    wb = xlrd.open_workbook(str(path))
    lines = []
    for sheet in wb.sheets():
        for i in range(sheet.nrows):
            row_text = " | ".join(
                str(c) for c in sheet.row_values(i) if str(c).strip()
            )
            if row_text.strip():
                lines.append(row_text)
    return "\n".join(lines)


def _extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _extract_pdf(path)
    if ext == ".docx":
        return _extract_docx(path)
    if ext == ".xlsx":
        return _extract_xlsx(path)
    if ext == ".xls":
        return _extract_xls(path)
    if ext == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")
    return ""


class LocalFolderKnowledgeService:
    """Reads and chunks all documents inside the EUZ Project docs folder."""

    def __init__(self, folder_path: str | None = None, chunk_size: int | None = None):
        raw = folder_path or settings.EUZ_DOCS_FOLDER
        p = Path(raw)
        self.folder_path = p if p.is_absolute() else _BACKEND_ROOT / p
        self.chunk_size = chunk_size or settings.GOOGLE_DRIVE_CHUNK_SIZE

    @property
    def configured(self) -> bool:
        return self.folder_path.exists() and self.folder_path.is_dir()

    def fetch_chunks(self) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
        if not self.configured:
            raise RuntimeError(
                f"EUZ docs folder not found: {self.folder_path}. "
                "Set EUZ_DOCS_FOLDER in rag_backend/.env."
            )

        chunks: List[Dict[str, Any]] = []
        files_ok = 0
        files_skipped = 0
        errors: List[Dict[str, str]] = []

        for path in sorted(self.folder_path.rglob("*")):
            if not path.is_file():
                continue
            if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
                files_skipped += 1
                continue

            try:
                text = _extract_text(path).strip()
            except Exception as exc:
                logger.warning("LocalFolder: failed to extract %s — %s", path.name, exc)
                errors.append({"file": path.name, "error": str(exc)})
                continue

            if not text:
                files_skipped += 1
                continue

            rel = path.relative_to(self.folder_path)
            topic = str(rel.parent) if str(rel.parent) != "." else path.stem
            tags = f"euz_project,{path.suffix.lstrip('.')}"
            file_key = hashlib.md5(str(rel).encode()).hexdigest()[:8]

            for idx, piece in enumerate(_chunk_text(text, self.chunk_size)):
                chunks.append(
                    {
                        "source": f"local_folder:{path.name}",
                        "chunk_id": f"{file_key}-{idx}",
                        "text": piece,
                        "topic": topic,
                        "tags": tags,
                    }
                )

            files_ok += 1

        stats: Dict[str, Any] = {
            "files_processed": files_ok,
            "files_skipped": files_skipped,
            "total_chunks": len(chunks),
            "errors": errors,
        }
        logger.info(
            "LocalFolder ingestion: %d files → %d chunks (%d skipped, %d errors)",
            files_ok,
            len(chunks),
            files_skipped,
            len(errors),
        )
        return chunks, stats
